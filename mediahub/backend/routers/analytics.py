from fastapi import APIRouter, HTTPException, Depends, Query
from fastapi.responses import StreamingResponse
from utils import get_db, get_current_user_id, require_group_member
from typing import Optional
from datetime import datetime, timedelta
from urllib.parse import quote
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
import io

router = APIRouter()


def _build_workbook(dt_start, dt_end, summary: dict, timeline, pl_stats, top_posts):
    BLUE = "1D4ED8"
    WHITE = "FFFFFF"
    GRAY = "F1F5F9"
    DARK = "1E293B"

    header_font = Font(bold=True, color=WHITE, size=11)
    header_fill = PatternFill("solid", fgColor=BLUE)
    header_align = Alignment(horizontal="center", vertical="center", wrap_text=True)
    cell_align = Alignment(vertical="center")
    thin_border = Border(
        bottom=Side(style="thin", color="CBD5E1"),
        right=Side(style="thin", color="E2E8F0"),
    )

    def style_header_row(ws, row, cols):
        for col in range(1, cols + 1):
            cell = ws.cell(row=row, column=col)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = header_align
            cell.border = thin_border

    def style_data_row(ws, row, cols, shade=False):
        fill = PatternFill("solid", fgColor=GRAY) if shade else None
        for col in range(1, cols + 1):
            cell = ws.cell(row=row, column=col)
            cell.alignment = cell_align
            cell.border = thin_border
            if shade:
                cell.fill = fill

    period_label = f"{dt_start.strftime('%d.%m.%Y')} — {dt_end.strftime('%d.%m.%Y')}"

    wb = openpyxl.Workbook()

    ws1 = wb.active
    ws1.title = "Сводка"
    ws1.row_dimensions[1].height = 30
    ws1["A1"] = f"Отчёт по аналитике — {period_label}"
    ws1["A1"].font = Font(bold=True, size=14, color=DARK)
    ws1["A1"].alignment = Alignment(horizontal="left", vertical="center")
    ws1.merge_cells("A1:B1")
    ws1.append([])

    headers = ["Показатель", "Значение"]
    ws1.append(headers)
    style_header_row(ws1, 3, 2)
    ws1.row_dimensions[3].height = 24

    rows_data = [
        ("Всего постов", summary["total"]),
        ("Опубликовано", summary["published"]),
        ("Запланировано", summary["scheduled"]),
        ("Черновиков", summary["drafts"]),
        ("Всего просмотров", summary["total_views"]),
        ("Реакции", summary["total_reactions"]),
        ("Комментарии", summary["total_comments"]),
        ("Репосты", summary["total_shares"]),
        ("Средние просмотры", summary["avg_views"]),
        ("Вовлечённость, %", summary["eng"]),
    ]
    for i, (label, val) in enumerate(rows_data, start=4):
        ws1.append([label, val])
        ws1.row_dimensions[i].height = 20
        style_data_row(ws1, i, 2, shade=(i % 2 == 0))
    ws1.column_dimensions["A"].width = 28
    ws1.column_dimensions["B"].width = 18

    ws2 = wb.create_sheet("Динамика")
    ws2.append(["Дата", "Просмотры", "Реакции", "Публикаций"])
    style_header_row(ws2, 1, 4)
    ws2.row_dimensions[1].height = 24
    for i, row in enumerate(timeline, start=2):
        ws2.append([row["date"], row["views"], row["reactions"], row["posts"]])
        ws2.row_dimensions[i].height = 18
        style_data_row(ws2, i, 4, shade=(i % 2 == 0))
    for col, w in zip("ABCD", [14, 14, 12, 14]):
        ws2.column_dimensions[col].width = w

    ws3 = wb.create_sheet("Площадки")
    ws3.append(["Площадка", "Публикаций", "Просмотры", "Реакции"])
    style_header_row(ws3, 1, 4)
    ws3.row_dimensions[1].height = 24
    for i, pl in enumerate(pl_stats, start=2):
        ws3.append([pl["platform"], pl["count"], pl["views"], pl["reactions"]])
        ws3.row_dimensions[i].height = 20
        style_data_row(ws3, i, 4, shade=(i % 2 == 0))
    for col, w in zip("ABCD", [16, 14, 14, 12]):
        ws3.column_dimensions[col].width = w

    ws4 = wb.create_sheet("Топ постов")
    ws4.append(["#", "Заголовок", "Просмотры", "Реакции", "Комментарии", "Репосты", "Дата публикации"])
    style_header_row(ws4, 1, 7)
    ws4.row_dimensions[1].height = 24
    for i, p in enumerate(top_posts, start=2):
        ws4.append([
            i - 1, p["title"], p["views"] or 0, p["reactions"] or 0,
            p["comments"] or 0, p["shares"] or 0, str(p["published_at"] or ""),
        ])
        ws4.row_dimensions[i].height = 20
        style_data_row(ws4, i, 7, shade=(i % 2 == 0))
    for col, w in zip("ABCDEFG", [4, 40, 12, 10, 14, 10, 22]):
        ws4.column_dimensions[col].width = w

    return wb


# ── Global Analytics ──────────────────────────────────────────────────────────

@router.get("/api/analytics/summary")
def analytics_summary():
    conn = get_db()
    c = conn.cursor()
    c.execute("SELECT COUNT(*) FROM posts")
    total = c.fetchone()["count"]
    c.execute("SELECT COUNT(*) FROM posts WHERE status='published'")
    published = c.fetchone()["count"]
    c.execute("SELECT COUNT(*) FROM posts WHERE status='scheduled'")
    scheduled = c.fetchone()["count"]
    c.execute("SELECT COUNT(*) FROM posts WHERE status='draft'")
    drafts = c.fetchone()["count"]
    c.execute("SELECT SUM(views) v,SUM(reactions) r,SUM(comments) c,SUM(shares) sh FROM posts WHERE status='published'")
    s = c.fetchone()
    c.execute(
        "SELECT id,title,views,reactions,comments,shares,published_at FROM posts "
        "WHERE status='published' ORDER BY (views+reactions*3+comments*2+shares*4) DESC LIMIT 5"
    )
    top = c.fetchall()
    pl_stats = []
    for pl in ["vk", "telegram"]:
        c.execute(
            "SELECT COUNT(*) cnt,SUM(views) v,SUM(reactions) r FROM posts "
            "WHERE platforms LIKE %s AND status='published'",
            (f'%"{pl}"%',),
        )
        ps = c.fetchone()
        pl_stats.append({"platform": pl, "count": ps["cnt"] or 0, "views": ps["v"] or 0, "reactions": ps["r"] or 0})
    conn.close()
    total_views = s["v"] or 0
    eng = round(((s["r"] or 0) + (s["c"] or 0)) / max(total_views, 1) * 100, 1)
    return {
        "total_posts": total, "published": published, "scheduled": scheduled, "drafts": drafts,
        "total_views": total_views, "total_reactions": s["r"] or 0,
        "total_comments": s["c"] or 0, "total_shares": s["sh"] or 0,
        "avg_views": round(total_views / max(published, 1)),
        "engagement_rate": eng,
        "top_posts": [dict(r) for r in top],
        "platform_stats": pl_stats,
    }


@router.get("/api/analytics/timeline")
def analytics_timeline(period: str = "month"):
    days = {"week": 7, "month": 30, "quarter": 90}.get(period, 30)
    now = datetime.now()
    result = []
    conn = get_db()
    c = conn.cursor()
    for i in range(days - 1, -1, -1):
        day = now - timedelta(days=i)
        ds = day.strftime("%Y-%m-%d")
        c.execute(
            "SELECT SUM(views) v, SUM(reactions) r, COUNT(*) p FROM posts WHERE published_at LIKE %s",
            (ds + "%",),
        )
        row = c.fetchone()
        result.append({
            "date": ds, "label": day.strftime("%d.%m"),
            "views": row["v"] or 0, "reactions": row["r"] or 0, "posts": row["p"] or 0,
        })
    conn.close()
    return result


@router.get("/api/analytics/export")
def analytics_export(start_date: str = Query(...), end_date: str = Query(...)):
    try:
        dt_start = datetime.strptime(start_date, "%Y-%m-%d")
        dt_end = datetime.strptime(end_date, "%Y-%m-%d")
    except ValueError:
        raise HTTPException(status_code=400, detail="Формат дат: YYYY-MM-DD")

    conn = get_db()
    c = conn.cursor()

    date_filter = "published_at >= %s AND published_at < %s AND status='published'"
    end_next = (dt_end + timedelta(days=1)).strftime("%Y-%m-%d")
    start_str = dt_start.strftime("%Y-%m-%d")

    c.execute(f"SELECT COUNT(*) FROM posts WHERE {date_filter}", (start_str, end_next))
    published = c.fetchone()["count"]
    c.execute("SELECT COUNT(*) FROM posts WHERE status='scheduled'")
    scheduled = c.fetchone()["count"]
    c.execute("SELECT COUNT(*) FROM posts WHERE status='draft'")
    drafts = c.fetchone()["count"]
    c.execute(f"SELECT COUNT(*) FROM posts WHERE published_at >= %s AND published_at < %s", (start_str, end_next))
    total = c.fetchone()["count"]
    c.execute(
        f"SELECT SUM(views) v,SUM(reactions) r,SUM(comments) cm,SUM(shares) sh FROM posts WHERE {date_filter}",
        (start_str, end_next),
    )
    s = c.fetchone()
    total_views = s["v"] or 0
    eng = round(((s["r"] or 0) + (s["cm"] or 0)) / max(total_views, 1) * 100, 1)

    timeline = []
    delta = (dt_end - dt_start).days + 1
    for i in range(delta):
        day = dt_start + timedelta(days=i)
        ds = day.strftime("%Y-%m-%d")
        c.execute("SELECT SUM(views) v, SUM(reactions) r, COUNT(*) p FROM posts WHERE published_at LIKE %s", (ds + "%",))
        row = c.fetchone()
        timeline.append({"date": ds, "label": day.strftime("%d.%m"), "views": row["v"] or 0, "reactions": row["r"] or 0, "posts": row["p"] or 0})

    pl_stats = []
    for pl in ["vk", "telegram"]:
        c.execute(
            f"SELECT COUNT(*) cnt,SUM(views) v,SUM(reactions) r FROM posts WHERE platforms LIKE %s AND {date_filter}",
            (f'%"{pl}"%', start_str, end_next),
        )
        ps = c.fetchone()
        pl_stats.append({"platform": pl.upper(), "count": ps["cnt"] or 0, "views": ps["v"] or 0, "reactions": ps["r"] or 0})

    c.execute(
        f"SELECT title,views,reactions,comments,shares,published_at FROM posts WHERE {date_filter} "
        "ORDER BY (views+reactions*3+comments*2+shares*4) DESC LIMIT 10",
        (start_str, end_next),
    )
    top_posts = c.fetchall()
    conn.close()

    summary = {
        "total": total, "published": published, "scheduled": scheduled, "drafts": drafts,
        "total_views": total_views, "total_reactions": s["r"] or 0,
        "total_comments": s["cm"] or 0, "total_shares": s["sh"] or 0,
        "avg_views": round(total_views / max(published, 1)), "eng": eng,
    }

    wb = _build_workbook(dt_start, dt_end, summary, timeline, pl_stats, top_posts)
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)

    month_ru = ["январь", "февраль", "март", "апрель", "май", "июнь",
                 "июль", "август", "сентябрь", "октябрь", "ноябрь", "декабрь"]
    fname = f"аналитика_{month_ru[dt_start.month - 1]}_{dt_start.year}.xlsx"
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename*=UTF-8''{quote(fname, safe='')}"},
    )


# ── Group-scoped Analytics ──────────────────────────────────────────────────────

@router.get("/api/groups/{gid}/analytics/summary")
def group_analytics_summary(gid: int, user_id: int = Depends(get_current_user_id)):
    conn = get_db()
    c = conn.cursor()
    require_group_member(gid, user_id, conn)
    c.execute("SELECT COUNT(*) FROM posts WHERE group_id=%s", (gid,))
    total = c.fetchone()["count"]
    c.execute("SELECT COUNT(*) FROM posts WHERE group_id=%s AND status='published'", (gid,))
    published = c.fetchone()["count"]
    c.execute("SELECT COUNT(*) FROM posts WHERE group_id=%s AND status='scheduled'", (gid,))
    scheduled = c.fetchone()["count"]
    c.execute("SELECT COUNT(*) FROM posts WHERE group_id=%s AND status='draft'", (gid,))
    drafts = c.fetchone()["count"]
    c.execute("SELECT SUM(views) v,SUM(reactions) r,SUM(comments) c,SUM(shares) sh FROM posts WHERE group_id=%s AND status='published'", (gid,))
    s = c.fetchone()
    c.execute(
        "SELECT id,title,views,reactions,comments,shares,published_at FROM posts "
        "WHERE group_id=%s AND status='published' ORDER BY (views+reactions*3+comments*2+shares*4) DESC LIMIT 5",
        (gid,),
    )
    top = c.fetchall()
    pl_stats = []
    for pl in ["vk", "telegram"]:
        c.execute(
            "SELECT COUNT(*) cnt,SUM(views) v,SUM(reactions) r FROM posts "
            "WHERE group_id=%s AND platforms LIKE %s AND status='published'",
            (gid, f'%"{pl}"%'),
        )
        ps = c.fetchone()
        pl_stats.append({"platform": pl, "count": ps["cnt"] or 0, "views": ps["v"] or 0, "reactions": ps["r"] or 0})
    conn.close()
    total_views = s["v"] or 0
    eng = round(((s["r"] or 0) + (s["c"] or 0)) / max(total_views, 1) * 100, 1)
    return {
        "total_posts": total, "published": published, "scheduled": scheduled, "drafts": drafts,
        "total_views": total_views, "total_reactions": s["r"] or 0,
        "total_comments": s["c"] or 0, "total_shares": s["sh"] or 0,
        "avg_views": round(total_views / max(published, 1)),
        "engagement_rate": eng,
        "top_posts": [dict(r) for r in top],
        "platform_stats": pl_stats,
    }


@router.get("/api/groups/{gid}/analytics/timeline")
def group_analytics_timeline(gid: int, period: str = "month", user_id: int = Depends(get_current_user_id)):
    days = {"week": 7, "month": 30, "quarter": 90}.get(period, 30)
    now = datetime.now()
    result = []
    conn = get_db()
    c = conn.cursor()
    require_group_member(gid, user_id, conn)
    for i in range(days - 1, -1, -1):
        day = now - timedelta(days=i)
        ds = day.strftime("%Y-%m-%d")
        c.execute(
            "SELECT SUM(views) v, SUM(reactions) r, COUNT(*) p FROM posts WHERE group_id=%s AND published_at LIKE %s",
            (gid, ds + "%"),
        )
        row = c.fetchone()
        result.append({
            "date": ds, "label": day.strftime("%d.%m"),
            "views": row["v"] or 0, "reactions": row["r"] or 0, "posts": row["p"] or 0,
        })
    conn.close()
    return result


@router.get("/api/groups/{gid}/analytics/export")
def group_analytics_export(
    gid: int,
    start_date: str = Query(...),
    end_date: str = Query(...),
    user_id: int = Depends(get_current_user_id),
):
    try:
        dt_start = datetime.strptime(start_date, "%Y-%m-%d")
        dt_end = datetime.strptime(end_date, "%Y-%m-%d")
    except ValueError:
        raise HTTPException(status_code=400, detail="Формат дат: YYYY-MM-DD")

    conn = get_db()
    c = conn.cursor()
    require_group_member(gid, user_id, conn)

    date_filter = "group_id=%s AND published_at >= %s AND published_at < %s AND status='published'"
    end_next = (dt_end + timedelta(days=1)).strftime("%Y-%m-%d")
    start_str = dt_start.strftime("%Y-%m-%d")

    c.execute(f"SELECT COUNT(*) FROM posts WHERE {date_filter}", (gid, start_str, end_next))
    published = c.fetchone()["count"]
    c.execute("SELECT COUNT(*) FROM posts WHERE group_id=%s AND status='scheduled'", (gid,))
    scheduled = c.fetchone()["count"]
    c.execute("SELECT COUNT(*) FROM posts WHERE group_id=%s AND status='draft'", (gid,))
    drafts = c.fetchone()["count"]
    c.execute(f"SELECT COUNT(*) FROM posts WHERE group_id=%s AND published_at >= %s AND published_at < %s", (gid, start_str, end_next))
    total = c.fetchone()["count"]
    c.execute(
        f"SELECT SUM(views) v,SUM(reactions) r,SUM(comments) cm,SUM(shares) sh FROM posts WHERE {date_filter}",
        (gid, start_str, end_next),
    )
    s = c.fetchone()
    total_views = s["v"] or 0
    eng = round(((s["r"] or 0) + (s["cm"] or 0)) / max(total_views, 1) * 100, 1)

    timeline = []
    delta = (dt_end - dt_start).days + 1
    for i in range(delta):
        day = dt_start + timedelta(days=i)
        ds = day.strftime("%Y-%m-%d")
        c.execute(
            "SELECT SUM(views) v, SUM(reactions) r, COUNT(*) p FROM posts WHERE group_id=%s AND published_at LIKE %s",
            (gid, ds + "%"),
        )
        row = c.fetchone()
        timeline.append({"date": ds, "label": day.strftime("%d.%m"), "views": row["v"] or 0, "reactions": row["r"] or 0, "posts": row["p"] or 0})

    pl_stats = []
    for pl in ["vk", "telegram"]:
        c.execute(
            f"SELECT COUNT(*) cnt,SUM(views) v,SUM(reactions) r FROM posts WHERE group_id=%s AND platforms LIKE %s AND {date_filter}",
            (gid, f'%"{pl}"%', start_str, end_next),
        )
        ps = c.fetchone()
        pl_stats.append({"platform": pl.upper(), "count": ps["cnt"] or 0, "views": ps["v"] or 0, "reactions": ps["r"] or 0})

    c.execute(
        f"SELECT title,views,reactions,comments,shares,published_at FROM posts WHERE {date_filter} "
        "ORDER BY (views+reactions*3+comments*2+shares*4) DESC LIMIT 10",
        (gid, start_str, end_next),
    )
    top_posts = c.fetchall()
    conn.close()

    summary = {
        "total": total, "published": published, "scheduled": scheduled, "drafts": drafts,
        "total_views": total_views, "total_reactions": s["r"] or 0,
        "total_comments": s["cm"] or 0, "total_shares": s["sh"] or 0,
        "avg_views": round(total_views / max(published, 1)), "eng": eng,
    }

    wb = _build_workbook(dt_start, dt_end, summary, timeline, pl_stats, top_posts)
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)

    month_ru = ["январь", "февраль", "март", "апрель", "май", "июнь",
                 "июль", "август", "сентябрь", "октябрь", "ноябрь", "декабрь"]
    fname = f"аналитика_{month_ru[dt_start.month - 1]}_{dt_start.year}.xlsx"
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename*=UTF-8''{quote(fname, safe='')}"},
    )
