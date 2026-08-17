"""
Медиатека и отчёт для учредителя.

Медиатека стережёт три вещи:

* **Оба источника видны.** Ради переиспользования всё и затевалось: раньше
  выбрать можно было только загрузки волонтёров, а снятое на прошлом
  мероприятии и уже вставленное в пост лежало в базе недостижимым.
* **Дублей нет.** Один файл легко попадает в оба источника разом — волонтёр
  прислал, редактор вставил в пост. Наружу он обязан уйти один раз.
* **Ссылки подписаны.** `/uploads` без подписи не отдаёт ничего, а заголовок
  Authorization к `<img src>` не приложить.

Отчёт стережёт то, ради чего он и другой документ: числа в нём не выдуманы, а
разбивка по направлениям существует.
"""
import uuid

import pytest
from conftest import auth

from utils import get_db


def make_group(client, token: str) -> int:
    r = client.post("/api/groups", json={"name": f"Группа {uuid.uuid4().hex[:6]}"},
                    headers=auth(token))
    assert r.status_code == 200, r.text
    return r.json()["id"]


def add_member(client, admin_token: str, gid: int, member_token: str, role: str) -> None:
    link = client.post(f"/api/groups/{gid}/invites", json={"role": role, "expires_hours": 24},
                       headers=auth(admin_token))
    assert link.status_code == 200, link.text
    r = client.post(f"/api/invites/{link.json()['token']}/accept", headers=auth(member_token))
    assert r.status_code == 200, r.text


def media(name: str, kind: str = "image") -> dict:
    return {"url": f"/uploads/{name}", "type": kind, "filename": name}


@pytest.fixture()
def library(client, make_user):
    """
    Группа, в которой есть и то и другое: одобренная загрузка волонтёра и пост
    со своими файлами. Один файл намеренно общий для обоих источников.
    """
    admin, _ = make_user("lib-admin")
    volunteer, _ = make_user("lib-vol")
    gid = make_group(client, admin)
    add_member(client, admin, gid, volunteer, "volunteer")

    shared = media("20260101120000_abc12345.jpg")
    only_volunteer = media("20260101120001_bcd23456.jpg")
    only_post = media("20260101120002_cde34567.jpg")

    up = client.post(
        f"/api/groups/{gid}/volunteer-media",
        json={"event_name": "Фестиваль «Весна»", "media": [shared, only_volunteer]},
        headers=auth(volunteer),
    )
    assert up.status_code == 200, up.text
    vid = up.json()["id"]
    r = client.put(f"/api/groups/{gid}/volunteer-media/{vid}/status",
                   json={"status": "approved"}, headers=auth(admin))
    assert r.status_code == 200, r.text

    p = client.post(
        f"/api/groups/{gid}/posts",
        json={"title": "Отчёт о фестивале", "content": "текст", "status": "draft",
              "media": [shared, only_post]},
        headers=auth(admin),
    )
    assert p.status_code == 200, p.text

    return {"gid": gid, "admin": admin, "volunteer": volunteer,
            "shared": shared, "only_volunteer": only_volunteer, "only_post": only_post}


def urls(items: list[dict]) -> set[str]:
    """Голые пути без подписи — сравнивать удобно именно их."""
    return {i["url"].split("?", 1)[0] for i in items}


# ── Медиатека ───────────────────────────────────────────────────────────────

def test_library_shows_both_sources(client, library):
    r = client.get(f"/api/groups/{library['gid']}/media-library",
                   headers=auth(library["admin"]))
    assert r.status_code == 200, r.text
    got = urls(r.json()["items"])
    assert library["only_volunteer"]["url"] in got, "загрузка волонтёра потерялась"
    assert library["only_post"]["url"] in got, "файл из поста недоступен для повторного использования"


def test_library_does_not_repeat_the_same_file(client, library):
    """
    Файл, попавший и в загрузку волонтёра, и в пост, — это один файл. Две
    карточки на него превратили бы медиатеку в склад дублей.
    """
    r = client.get(f"/api/groups/{library['gid']}/media-library",
                   headers=auth(library["admin"]))
    items = r.json()["items"]
    paths = [i["url"].split("?", 1)[0] for i in items]
    assert len(paths) == len(set(paths)), f"дубли в медиатеке: {paths}"


def test_library_links_are_signed(client, library):
    r = client.get(f"/api/groups/{library['gid']}/media-library",
                   headers=auth(library["admin"]))
    for item in r.json()["items"]:
        assert "sig=" in item["url"] and "exp=" in item["url"], item["url"]
    # И подпись настоящая, а не любая строка в параметре: файлов на диске в
    # тесте нет, поэтому проверяем саму проверку подписи, а не выдачу файла.
    from urllib.parse import parse_qs, urlparse

    from utils import check_upload_signature, upload_filename

    url = r.json()["items"][0]["url"]
    query = parse_qs(urlparse(url).query)
    check_upload_signature(upload_filename(url), query["exp"][0], query["sig"][0])


def test_library_hides_unapproved_volunteer_uploads(client, make_user):
    """Модерация затем и нужна: неодобренное в редактор попадать не должно."""
    admin, _ = make_user("mod-admin")
    volunteer, _ = make_user("mod-vol")
    gid = make_group(client, admin)
    add_member(client, admin, gid, volunteer, "volunteer")
    pending = media("20260101120003_def45678.jpg")
    client.post(f"/api/groups/{gid}/volunteer-media",
                json={"event_name": "Ещё не смотрели", "media": [pending]},
                headers=auth(volunteer))

    r = client.get(f"/api/groups/{gid}/media-library", headers=auth(admin))
    assert pending["url"] not in urls(r.json()["items"])


def test_library_search_looks_at_names_and_events(client, library):
    """Человек помнит «это было с фестиваля», а не имя файла."""
    r = client.get(f"/api/groups/{library['gid']}/media-library?q=фестиваль",
                   headers=auth(library["admin"]))
    assert r.status_code == 200, r.text
    assert r.json()["items"], "поиск по названию мероприятия ничего не нашёл"
    assert r.json()["total"] == len(r.json()["items"])


def test_library_filters_by_source(client, library):
    r = client.get(f"/api/groups/{library['gid']}/media-library?source=post",
                   headers=auth(library["admin"]))
    assert all(i["source"] == "post" for i in r.json()["items"])
    assert library["only_volunteer"]["url"] not in urls(r.json()["items"])


def test_library_total_matches_the_filter(client, library):
    """
    Счётчик обязан считать то же, что показываем: «показаны 1–2 из 40» при
    двух найденных — это неверная подпись и лишние страницы в навигации.
    """
    r = client.get(f"/api/groups/{library['gid']}/media-library?source=volunteer",
                   headers=auth(library["admin"]))
    data = r.json()
    assert data["total"] == len(data["items"])


def test_library_is_closed_to_watchers(client, library):
    r = client.get(f"/api/groups/{library['gid']}/media-library",
                   headers=auth(library["volunteer"]))
    assert r.status_code == 403


def test_library_is_closed_to_outsiders(client, library, make_user):
    stranger, _ = make_user("stranger")
    r = client.get(f"/api/groups/{library['gid']}/media-library", headers=auth(stranger))
    assert r.status_code == 403


def test_volunteer_media_total_respects_filters(client, library):
    """
    Счётчик в списке загрузок брал всю группу целиком, не глядя на фильтры и на
    то, что волонтёр видит только свои: подпись «из 120» при трёх видимых.
    """
    r = client.get(f"/api/groups/{library['gid']}/volunteer-media?status=rejected",
                   headers=auth(library["admin"]))
    data = r.json()
    assert data["total"] == len(data["items"]) == 0


# ── Отчёт для учредителя ────────────────────────────────────────────────────

XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def published_post(client, token: str, gid: int, title: str, tags: list[str]) -> int:
    r = client.post(
        f"/api/groups/{gid}/posts",
        json={"title": title, "content": "текст", "status": "draft",
              "tags": tags, "platforms": ["vk"]},
        headers=auth(token),
    )
    assert r.status_code == 200, r.text
    pid = r.json()["id"]
    conn = get_db()
    c = conn.cursor()
    c.execute("UPDATE posts SET status='published', published_at=NOW() WHERE id=%s", (pid,))
    conn.commit()
    conn.close()
    return pid


def test_report_downloads_as_a_spreadsheet(client, make_user):
    token, _ = make_user("report")
    gid = make_group(client, token)
    published_post(client, token, gid, "Открытие сезона", ["мероприятия"])

    r = client.get(f"/api/groups/{gid}/analytics/report?start_date=2020-01-01&end_date=2030-12-31",
                   headers=auth(token))
    assert r.status_code == 200, r.text
    assert r.headers["content-type"] == XLSX
    # xlsx — это zip, и начинается он с PK.
    assert r.content[:2] == b"PK"
    assert len(r.content) > 4000, "подозрительно маленький файл"


def test_report_counts_by_direction(client, make_user):
    """
    Разбивка по направлениям — то, ради чего учредитель отчёт и открывает:
    сколько было про мероприятия, сколько про гранты.
    """
    import io

    import openpyxl

    token, _ = make_user("report-tags")
    gid = make_group(client, token)
    published_post(client, token, gid, "Фестиваль", ["мероприятия"])
    published_post(client, token, gid, "Субботник", ["мероприятия"])
    published_post(client, token, gid, "Конкурс грантов", ["гранты"])

    r = client.get(f"/api/groups/{gid}/analytics/report?start_date=2020-01-01&end_date=2030-12-31",
                   headers=auth(token))
    wb = openpyxl.load_workbook(io.BytesIO(r.content))
    assert "Направления" in wb.sheetnames
    rows = {row[0]: row[1] for row in wb["Направления"].iter_rows(min_row=2, values_only=True)}
    assert rows.get("мероприятия") == 2
    assert rows.get("гранты") == 1


def test_report_says_no_data_instead_of_inventing_zero(client, make_user):
    """
    Просмотры Telegram недостижимы в принципе, ВК могли не синхронизировать.
    Пустая ячейка читается как ноль, поэтому пишем словами.
    """
    import io

    import openpyxl

    token, _ = make_user("report-nodata")
    gid = make_group(client, token)
    published_post(client, token, gid, "Без статистики", ["новости"])

    r = client.get(f"/api/groups/{gid}/analytics/report?start_date=2020-01-01&end_date=2030-12-31",
                   headers=auth(token))
    wb = openpyxl.load_workbook(io.BytesIO(r.content))
    values = [c.value for row in wb["Отчёт"].iter_rows() for c in row if c.value]
    assert any(str(v) == "нет данных" for v in values), \
        "по несобранным просмотрам отчёт должен говорить «нет данных», а не показывать ноль"
    # И ни один показатель не подменён нулём: ноль в отчёте наружу читается как
    # «людям не зашло», хотя на деле цифр просто нет.
    assert not any(v == 0 for v in values), f"в отчёте есть выдуманные нули: {values}"


def test_report_refuses_a_backwards_period(client, make_user):
    token, _ = make_user("report-bad")
    gid = make_group(client, token)
    r = client.get(f"/api/groups/{gid}/analytics/report?start_date=2026-12-31&end_date=2026-01-01",
                   headers=auth(token))
    assert r.status_code == 400


def test_report_is_closed_to_outsiders(client, make_user):
    owner, _ = make_user("report-owner")
    stranger, _ = make_user("report-stranger")
    gid = make_group(client, owner)
    r = client.get(f"/api/groups/{gid}/analytics/report?start_date=2026-01-01&end_date=2026-12-31",
                   headers=auth(stranger))
    assert r.status_code == 403
