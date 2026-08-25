import io
from datetime import datetime, timedelta
from urllib.parse import quote

import openpyxl
from fastapi import APIRouter, Depends, HTTPException, Query
from fastapi.responses import StreamingResponse
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

import comments
from utils import (
    app_now,
    as_json_list,
    get_current_user_id,
    get_db,
    posts_scope,
    require_group_member,
)

router = APIRouter()

# Статистика лежит по площадкам отдельно (post_stats), поэтому «есть ли цифры»
# больше не захардкожено, а видно по данным: если по площадке нет ни одной
# собранной строки, отдаём None («нет данных»), а не 0 и не чужие числа. Как
# только сбор по площадке появится, признак переключится сам.
def _platform_row(platform: str, count: int, collected: int, views, reactions,
                  label: str | None = None) -> dict:
    known = bool(collected)
    return {
        "platform": label or platform,
        "count": count or 0,
        "collected": collected or 0,
        "views": (views or 0) if known else None,
        "reactions": (reactions or 0) if known else None,
        "stats_available": known,
    }


# «Нет данных» в отчёте: None выглядел бы в ячейке пустым местом, а пустое
# место читается как ноль.
NO_DATA = "нет данных"


def _collected_views(row) -> int | None:
    """
    Сумма просмотров или None, если их вообще не собирали.

    Само по себе SUM(views) ноль не отличает: в post_totals суммы обёрнуты
    в COALESCE, а по площадкам вроде Telegram просмотры недостижимы в принципе.
    Различие несёт views_samples — COUNT по ненулевым значениям (миграция
    b2e6a3d95f41): ноль там означает ровно «цифр нет ни по одному посту».
    """
    if not row["vs"]:
        return None
    return row["v"] or 0


def _engagement(views, reactions, comments) -> float | None:
    """
    Вовлечённость — доля откликов от просмотров, или None, если считать не из чего.

    Знаменатель раньше страховали через `max(views, 1)`. При нуле просмотров и трёх
    реакциях это давало 300 %, а при нуле и нуле — аккуратные «0,0 %», которые
    читаются как «людям не заходит», хотя на деле цифр просто нет: у Telegram
    просмотры недостижимы в принципе, у ВК их могли ещё не синхронизировать.
    NULL в сумме просмотров означает «ни по одной площадке не собирали» — отдаём
    то же None, что и разбивка по площадкам, а интерфейс пишет «нет данных».
    """
    if not views:
        return None
    return round(((reactions or 0) + (comments or 0)) / views * 100, 1)


def _avg_views(views, published: int) -> int | None:
    """Средние просмотры на пост — по той же причине None, пока просмотров нет."""
    if views is None or published <= 0:
        return None
    return round(views / published)


# Сумма по всем площадкам живёт в представлении post_totals (см. миграцию
# f4b2e8c15d93) — иначе одинаковый подзапрос разъехался бы по полутора десяткам мест.
TOTALS = "posts JOIN post_totals t ON t.post_id = posts.id"


def _platform_stats(c, where: str, params: list, label_upper: bool = False) -> list[dict]:
    """Разбивка по площадкам: сколько постов ушло и что по ним собрано."""
    rows = []
    for pl in ("vk", "telegram"):
        c.execute(
            f"SELECT COUNT(*) cnt FROM posts WHERE platforms ? %s AND {where}",
            [pl] + list(params),
        )
        count = c.fetchone()["cnt"]
        c.execute(
            "SELECT COUNT(*) collected, SUM(s.views) v, SUM(s.reactions) r "
            "FROM post_stats s JOIN posts ON posts.id = s.post_id "
            f"WHERE s.platform=%s AND {where}",
            [pl] + list(params),
        )
        got = c.fetchone()
        rows.append(_platform_row(pl, count, got["collected"], got["v"], got["r"],
                                  label=pl.upper() if label_upper else None))
    return rows


def _build_workbook(dt_start, dt_end, summary: dict, timeline, pl_stats, top_posts):
    BLUE = "1D4ED8"
    WHITE = "FFFFFF"
    GRAY = "F1F5F9"
    DARK = "1E293B"

    header_font = Font(bold=True, color=WHITE, size=11)
    header_fill = PatternFill("solid", fgColor=BLUE)
    header_align = Alignment(horizontal="center", vertical="center", wrap_text=True)
    cell_align = Alignment(vertical="center")
    thin_border = Border(
        bottom=Side(style="thin", color="CBD5E1"),
        right=Side(style="thin", color="E2E8F0"),
    )

    def style_header_row(ws, row, cols):
        for col in range(1, cols + 1):
            cell = ws.cell(row=row, column=col)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = header_align
            cell.border = thin_border

    def style_data_row(ws, row, cols, shade=False):
        fill = PatternFill("solid", fgColor=GRAY) if shade else None
        for col in range(1, cols + 1):
            cell = ws.cell(row=row, column=col)
            cell.alignment = cell_align
            cell.border = thin_border
            if shade:
                cell.fill = fill

    period_label = f"{dt_start.strftime('%d.%m.%Y')} — {dt_end.strftime('%d.%m.%Y')}"

    wb = openpyxl.Workbook()

    ws1 = wb.active
    ws1.title = "Сводка"
    ws1.row_dimensions[1].height = 30
    ws1["A1"] = f"Отчёт по аналитике — {period_label}"
    ws1["A1"].font = Font(bold=True, size=14, color=DARK)
    ws1["A1"].alignment = Alignment(horizontal="left", vertical="center")
    ws1.merge_cells("A1:B1")
    ws1.append([])

    headers = ["Показатель", "Значение"]
    ws1.append(headers)
    style_header_row(ws1, 3, 2)
    ws1.row_dimensions[3].height = 24

    rows_data = [
        ("Всего постов", summary["total"]),
        ("Опубликовано", summary["published"]),
        ("Запланировано", summary["scheduled"]),
        ("Черновиков", summary["drafts"]),
        ("Всего просмотров", summary["total_views"]),
        ("Реакции", summary["total_reactions"]),
        ("Комментарии", summary["total_comments"]),
        ("Репосты", summary["total_shares"]),
        # Пока просмотры не собраны, среднего и вовлечённости не существует —
        # пустая ячейка прочиталась бы как ноль, поэтому пишем словами.
        ("Средние просмотры", summary["avg_views"] if summary["avg_views"] is not None else NO_DATA),
        ("Вовлечённость, %", summary["eng"] if summary["eng"] is not None else NO_DATA),
    ]
    for i, (label, val) in enumerate(rows_data, start=4):
        ws1.append([label, val])
        ws1.row_dimensions[i].height = 20
        style_data_row(ws1, i, 2, shade=(i % 2 == 0))
    ws1.column_dimensions["A"].width = 28
    ws1.column_dimensions["B"].width = 18

    ws2 = wb.create_sheet("Динамика")
    ws2.append(["Дата", "Просмотры", "Реакции", "Публикаций"])
    style_header_row(ws2, 1, 4)
    ws2.row_dimensions[1].height = 24
    for i, row in enumerate(timeline, start=2):
        ws2.append([row["date"], row["views"], row["reactions"], row["posts"]])
        ws2.row_dimensions[i].height = 18
        style_data_row(ws2, i, 4, shade=(i % 2 == 0))
    for col, w in zip("ABCD", [14, 14, 12, 14], strict=False):
        ws2.column_dimensions[col].width = w

    ws3 = wb.create_sheet("Площадки")
    ws3.append(["Площадка", "Публикаций", "Просмотры", "Реакции"])
    style_header_row(ws3, 1, 4)
    ws3.row_dimensions[1].height = 24
    for i, pl in enumerate(pl_stats, start=2):
        # None → пустая ячейка выглядит как ноль; пишем словами, что данных нет
        views = pl["views"] if pl["views"] is not None else NO_DATA
        reactions = pl["reactions"] if pl["reactions"] is not None else NO_DATA
        ws3.append([pl["platform"], pl["count"], views, reactions])
        ws3.row_dimensions[i].height = 20
        style_data_row(ws3, i, 4, shade=(i % 2 == 0))
    for col, w in zip("ABCD", [16, 14, 14, 12], strict=False):
        ws3.column_dimensions[col].width = w

    ws4 = wb.create_sheet("Топ постов")
    ws4.append(["#", "Заголовок", "Просмотры", "Реакции", "Комментарии", "Репосты", "Дата публикации"])
    style_header_row(ws4, 1, 7)
    ws4.row_dimensions[1].height = 24
    for i, p in enumerate(top_posts, start=2):
        ws4.append([
            i - 1, p["title"], p["views"] or 0, p["reactions"] or 0,
            p["comments"] or 0, p["shares"] or 0, str(p["published_at"] or ""),
        ])
        ws4.row_dimensions[i].height = 20
        style_data_row(ws4, i, 7, shade=(i % 2 == 0))
    for col, w in zip("ABCDEFG", [4, 40, 12, 10, 14, 10, 22], strict=False):
        ws4.column_dimensions[col].width = w

    return wb


# ── Отчёт для учредителя ────────────────────────────────────────────────────
# Аналитика в интерфейсе отвечает на вопросы того, кто ведёт каналы: что зашло,
# когда публиковать, где просело. Учредителю нужен другой документ и другой
# язык — «что делали и что из этого вышло» связным текстом, а не сетка
# показателей. Отсюда отдельный отчёт, а не ещё одна вкладка в экспорте.


def _plural(n: int, one: str, few: str, many: str) -> str:
    """Русское согласование числительного: 1 публикация, 2 публикации, 5 публикаций."""
    if 11 <= n % 100 <= 14:
        return many
    last = n % 10
    if last == 1:
        return one
    if 2 <= last <= 4:
        return few
    return many


def _spaced(n: int) -> str:
    """Число с неразрывными пробелами по разрядам: 12 480, а не 12480."""
    return f"{n:,}".replace(",", " ")


def _tag_breakdown(c, where: str, params: list) -> list[dict]:
    """
    Разбивка по рубрикам. Учредителю интереснее всего именно она: сколько было
    про мероприятия, сколько про гранты, сколько про набор волонтёров.
    """
    c.execute(
        "SELECT tag, COUNT(*) AS posts, SUM(t.views) AS v, SUM(t.views_samples) AS vs, "
        "       SUM(t.reactions) AS r "
        f"FROM {TOTALS} "
        "CROSS JOIN LATERAL jsonb_array_elements_text("
        "    CASE WHEN jsonb_typeof(posts.tags) = 'array' THEN posts.tags ELSE '[]'::jsonb END"
        ") AS tag "
        f"WHERE {where} "
        "GROUP BY tag ORDER BY posts DESC, tag",
        params,
    )
    return [
        {"tag": r["tag"], "posts": r["posts"],
         "views": _collected_views({"v": r["v"], "vs": r["vs"]}),
         "reactions": r["r"] or 0}
        for r in c.fetchall()
    ]


def _vk_community(c, gid: int | None) -> str | None:
    """Идентификатор сообщества ВК — из него собирается ссылка на запись."""
    if gid is None:
        c.execute("SELECT group_id FROM vk_settings WHERE id=1")
    else:
        c.execute("SELECT group_id FROM vk_settings WHERE workspace_id=%s", (gid,))
    row = c.fetchone()
    return row["group_id"] if row else None


def _post_link(post: dict, community: str | None) -> str:
    """
    Ссылка на запись во ВКонтакте. Учредителю нужен не наш идентификатор поста,
    а возможность открыть публикацию и посмотреть на неё.
    """
    if not post.get("vk_post_id") or not community:
        return ""
    return f"https://vk.com/wall-{str(community).lstrip('-')}_{post['vk_post_id']}"


def _report_summary_text(org: str, period: str, s: dict, published: int,
                         platforms: list[dict], top_tag: str | None,
                         replies: dict | None = None) -> list[str]:
    """Связный текст вместо таблицы: с этого учредитель начинает читать."""
    lines = [
        # Кавычки вокруг названия не ставим: названия учреждений сплошь и рядом
        # уже написаны в кавычках — «Молодёжный центр «Спектр»» читается плохо.
        f"За период {period} организация {org} опубликовала "
        f"{_spaced(published)} {_plural(published, 'материал', 'материала', 'материалов')} "
        f"в социальных сетях."
    ]

    where = [f"{p['platform']} — {_spaced(p['count'])}" for p in platforms if p["count"]]
    if where:
        lines.append("Распределение по площадкам: " + ", ".join(where) + ".")

    views = s["views"]
    if views is None:
        # Придумывать охват нельзя: у Telegram просмотры недостижимы в принципе,
        # у ВК их могли ещё не синхронизировать. Молчание честнее нуля.
        lines.append(
            "Данные о просмотрах за период не собраны, поэтому охват в отчёте не приводится."
        )
    else:
        lines.append(
            f"Публикации набрали {_spaced(views)} "
            f"{_plural(views, 'просмотр', 'просмотра', 'просмотров')}"
            + (f", в среднем {_spaced(s['avg_views'])} на материал"
               if s["avg_views"] is not None else "")
            + "."
        )

    parts = []
    for value, forms in (
        (s["reactions"], ('отметка «нравится»', 'отметки «нравится»', 'отметок «нравится»')),
        (s["comments"], ('комментарий', 'комментария', 'комментариев')),
        (s["shares"], ('репост', 'репоста', 'репостов')),
    ):
        # None — не собирали. Такой показатель в предложение не попадает вовсе:
        # написать «0 комментариев» там, где их не считали, значит соврать.
        if value:
            parts.append(f"{_spaced(value)} {_plural(value, *forms)}")
    if parts:
        total_reactions = sum(v for v in (s["reactions"], s["comments"], s["shares"]) if v)
        lines.append(
            f"Читатели откликнулись {_spaced(total_reactions)} "
            f"{_plural(total_reactions, 'раз', 'раза', 'раз')}: " + ", ".join(parts) + "."
        )

    if top_tag:
        lines.append(f"Больше всего материалов вышло по направлению «{top_tag}».")

    if replies and replies["total"]:
        answered = replies["answered"]
        line = (f"На публикации поступило {_spaced(replies['total'])} "
                f"{_plural(replies['total'], 'обращение', 'обращения', 'обращений')} "
                f"от жителей, отвечено на {_spaced(answered)}")
        if replies["avg_reply_hours"] is not None:
            hours = replies["avg_reply_hours"]
            line += (f"; среднее время ответа — {hours:.1f}".replace(".", ",") + " ч")
        lines.append(line + ".")

    return lines


def _build_report(org: str, dt_start, dt_end, s: dict, published: int,
                  platforms: list[dict], tags: list[dict], posts: list[dict],
                  community: str | None, replies: dict | None = None):
    INK = "101014"
    MUTED = "6B7280"
    RULE = "D4D4D8"

    period = f"с {dt_start.strftime('%d.%m.%Y')} по {dt_end.strftime('%d.%m.%Y')}"

    wb = openpyxl.Workbook()

    # ── Лист 1: собственно отчёт ────────────────────────────────────────────
    ws = wb.active
    ws.title = "Отчёт"
    ws.column_dimensions["A"].width = 34
    ws.column_dimensions["B"].width = 22
    ws.column_dimensions["C"].width = 18
    ws.column_dimensions["D"].width = 18
    ws.sheet_view.showGridLines = False

    ws["A1"] = "Отчёт об информационной работе"
    ws["A1"].font = Font(bold=True, size=16, color=INK)
    ws.row_dimensions[1].height = 26
    ws["A2"] = f"{org} · {period}"
    ws["A2"].font = Font(size=11, color=MUTED)
    ws.row_dimensions[2].height = 20

    row = 4
    for line in _report_summary_text(org, period, s, published, platforms,
                                     tags[0]["tag"] if tags else None, replies):
        ws.cell(row=row, column=1, value=line)
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=4)
        ws.cell(row=row, column=1).alignment = Alignment(wrap_text=True, vertical="top")
        # Высота под перенос: openpyxl автоподбор не умеет, а обрезанный текст
        # в отчёте наружу — это отчёт, который никто не дочитает.
        ws.row_dimensions[row].height = 15 * (1 + len(line) // 95)
        row += 1

    row += 1
    ws.cell(row=row, column=1, value="Ключевые показатели").font = Font(bold=True, size=12, color=INK)
    row += 1

    figures = [
        ("Опубликовано материалов", _spaced(published)),
        ("Просмотры", _spaced(s["views"]) if s["views"] is not None else NO_DATA),
        ("Среднее число просмотров на материал",
         _spaced(s["avg_views"]) if s["avg_views"] is not None else NO_DATA),
        ("Отметки «нравится»", _spaced(s["reactions"]) if s["reactions"] is not None else NO_DATA),
        ("Комментарии", _spaced(s["comments"]) if s["comments"] is not None else NO_DATA),
        ("Репосты", _spaced(s["shares"]) if s["shares"] is not None else NO_DATA),
        ("Отклик на просмотр",
         f"{s['engagement']:.1f} %".replace(".", ",") if s["engagement"] is not None else NO_DATA),
    ]
    if replies and replies["total"]:
        figures += [
            ("Обращений от жителей", _spaced(replies["total"])),
            ("Из них отвечено", _spaced(replies["answered"])),
            ("Среднее время ответа, ч",
             f"{replies['avg_reply_hours']:.1f}".replace(".", ",")
             if replies["avg_reply_hours"] is not None else NO_DATA),
        ]
    for label, value in figures:
        ws.cell(row=row, column=1, value=label).font = Font(size=11, color=INK)
        cell = ws.cell(row=row, column=2, value=value)
        cell.font = Font(bold=True, size=11, color=INK)
        cell.alignment = Alignment(horizontal="right")
        for col in (1, 2):
            ws.cell(row=row, column=col).border = Border(
                bottom=Side(style="thin", color=RULE))
        ws.row_dimensions[row].height = 19
        row += 1

    if any(s[k] is None for k in ("views", "reactions", "comments", "shares")):
        row += 1
        note = ws.cell(
            row=row, column=1,
            value="«нет данных» означает, что показатель не собирался, а не что он равен нулю. "
                  "Просмотры доступны только для ВКонтакте: Telegram их не отдаёт.",
        )
        note.font = Font(size=9, color=MUTED, italic=True)
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=4)
        note.alignment = Alignment(wrap_text=True, vertical="top")
        ws.row_dimensions[row].height = 28

    # ── Лист 2: направления работы ──────────────────────────────────────────
    def sheet(title: str, headers: list[str], widths: list[int]):
        s2 = wb.create_sheet(title)
        s2.sheet_view.showGridLines = False
        s2.append(headers)
        for i in range(1, len(headers) + 1):
            cell = s2.cell(row=1, column=i)
            cell.font = Font(bold=True, size=10, color=INK)
            cell.border = Border(bottom=Side(style="medium", color=INK))
            cell.alignment = Alignment(vertical="center")
        s2.row_dimensions[1].height = 22
        for col, w in zip(
            [chr(ord("A") + i) for i in range(len(widths))], widths, strict=True
        ):
            s2.column_dimensions[col].width = w
        return s2

    # Формат ячейки, а не форматирование строкой: число остаётся числом —
    # его можно сложить и по нему можно отсортировать, — но читается с
    # разрядами. Строка «20 280» не умеет ни того, ни другого.
    NUM = "# ##0"

    def numbers(sheet_obj, first_col: int, last_col: int) -> None:
        for row in sheet_obj.iter_rows(min_row=2, min_col=first_col, max_col=last_col):
            for cell in row:
                if isinstance(cell.value, int | float):
                    cell.number_format = NUM

    ws2 = sheet("Направления", ["Направление", "Материалов", "Просмотры", "Отклики"],
                [30, 14, 16, 14])
    for i, t in enumerate(tags, start=2):
        ws2.append([t["tag"], t["posts"],
                    t["views"] if t["views"] is not None else NO_DATA, t["reactions"]])
        ws2.cell(row=i, column=1).border = Border(bottom=Side(style="thin", color=RULE))
    numbers(ws2, 2, 4)
    if not tags:
        ws2.append(["Материалы за период не размечены по направлениям"])

    ws3 = sheet("Площадки", ["Площадка", "Материалов", "Просмотры", "Отклики"],
                [22, 14, 16, 14])
    for pl in platforms:
        ws3.append([pl["platform"], pl["count"],
                    pl["views"] if pl["views"] is not None else NO_DATA,
                    pl["reactions"] if pl["reactions"] is not None else NO_DATA])
    numbers(ws3, 2, 4)

    ws4 = sheet("Публикации",
                ["Дата", "Заголовок", "Направления", "Площадки", "Просмотры", "Отклики", "Ссылка"],
                [12, 46, 24, 16, 12, 12, 42])
    for i, p in enumerate(posts, start=2):
        link = _post_link(p, community)
        # ДД.ММ.ГГГГ, а не ISO: документ читает человек, а не программа.
        published_on = p["published_at"].strftime("%d.%m.%Y") if p["published_at"] else ""
        ws4.append([
            published_on,
            p["title"],
            ", ".join(as_json_list(p["tags"])),
            ", ".join("ВКонтакте" if x == "vk" else "Telegram" if x == "telegram" else x
                      for x in as_json_list(p["platforms"])),
            p["views"] if p["views_samples"] else NO_DATA,
            # Отклики — сумма только по тому, что действительно считали.
            ((p["reactions"] or 0) + (p["comments"] or 0) + (p["shares"] or 0))
            if (p["reactions_samples"] or p["comments_samples"] or p["shares_samples"])
            else NO_DATA,
            link,
        ])
        if link:
            cell = ws4.cell(row=i, column=7)
            cell.hyperlink = link
            cell.font = Font(color="1D4ED8", underline="single", size=10)
        ws4.cell(row=i, column=2).alignment = Alignment(wrap_text=False)
    numbers(ws4, 5, 6)
    if not posts:
        ws4.append(["", "За указанный период публикаций не было"])

    return wb


def _report(c, where: str, params: list, org: str, start_date: str, end_date: str,
            gid: int | None):
    try:
        dt_start = datetime.strptime(start_date, "%Y-%m-%d")
        dt_end = datetime.strptime(end_date, "%Y-%m-%d")
    except ValueError:
        raise HTTPException(status_code=400, detail="Формат дат: YYYY-MM-DD")
    if dt_end < dt_start:
        raise HTTPException(status_code=400, detail="Конец периода раньше начала")

    end_next = (dt_end + timedelta(days=1)).strftime("%Y-%m-%d")
    period = f"{where} AND published_at >= %s AND published_at < %s AND status='published'"
    period_params = params + [dt_start.strftime("%Y-%m-%d"), end_next]

    c.execute(f"SELECT COUNT(*) FROM posts WHERE {period}", period_params)
    published = c.fetchone()["count"]

    # Счётчики собранных значений тянем по каждому показателю, а не только по
    # просмотрам. Ноль реакций в отчёте наружу означает «людям не зашло» — а
    # если реакции просто не собирали, это неправда о работе организации,
    # напечатанная на бумаге и отданная учредителю.
    c.execute(
        f"SELECT SUM(t.views) v, SUM(t.views_samples) vs, "
        f"SUM(t.reactions) r, SUM(t.reactions_samples) rs, "
        f"SUM(t.comments) cm, SUM(t.comments_samples) cms, "
        f"SUM(t.shares) sh, SUM(t.shares_samples) shs "
        f"FROM {TOTALS} WHERE {period}",
        period_params,
    )
    row = c.fetchone()
    views = _collected_views(row)

    def collected(value_key: str, samples_key: str) -> int | None:
        return (row[value_key] or 0) if row[samples_key] else None

    summary = {
        "views": views,
        "reactions": collected("r", "rs"),
        "comments": collected("cm", "cms"),
        "shares": collected("sh", "shs"),
        "avg_views": _avg_views(views, published),
        "engagement": _engagement(views, row["r"], row["cm"]),
    }

    # Обращения граждан: для учреждения это не «ещё одна метрика», а предмет
    # отдельного спроса — отвечать в комментариях оно обязано.
    replies = comments.stats(c.connection, gid, dt_start, dt_end + timedelta(days=1)) \
        if gid is not None else None

    platforms = _platform_stats(c, period, period_params)
    for pl in platforms:
        pl["platform"] = "ВКонтакте" if pl["platform"] == "vk" else "Telegram"

    tags = _tag_breakdown(c, period, period_params)

    c.execute(
        f"SELECT title, published_at, tags, platforms, vk_post_id, "
        f"t.views, t.views_samples, t.reactions, t.comments, t.shares, "
        f"t.reactions_samples, t.comments_samples, t.shares_samples "
        f"FROM {TOTALS} WHERE {period} ORDER BY published_at",
        period_params,
    )
    posts = [dict(r) for r in c.fetchall()]

    wb = _build_report(org, dt_start, dt_end, summary, published, platforms, tags,
                       posts, _vk_community(c, gid), replies)
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)

    fname = (f"отчёт_{org}_{dt_start.strftime('%d.%m.%Y')}-"
             f"{dt_end.strftime('%d.%m.%Y')}.xlsx")
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename*=UTF-8''{quote(fname, safe='')}"},
    )


# ── Ядро: одно на глобальные и групповые ручки ──────────────────────────────
# Раньше это была пара почти дословных копий на каждый отчёт: глобальную и
# групповую отличало только условие выборки. Копии неизбежно разъезжаются —
# правку вносят в одну и забывают про вторую. Теперь различие сведено к паре
# (where, params), а логика одна.


def _summary(c, where: str, params: list) -> dict:
    def count(extra: str, extra_params: list | None = None) -> int:
        c.execute(f"SELECT COUNT(*) FROM posts WHERE {where}{extra}", params + (extra_params or []))
        return c.fetchone()["count"]

    total = count("")
    published = count(" AND status='published'")
    scheduled = count(" AND status='scheduled'")
    drafts = count(" AND status='draft'")
    on_review = count(" AND status='on_review'")

    c.execute(
        f"SELECT SUM(t.views) v,SUM(t.reactions) r,SUM(t.comments) cm,SUM(t.shares) sh,"
        f"SUM(t.views_samples) vs "
        f"FROM {TOTALS} WHERE {where} AND status='published'",
        params,
    )
    s = c.fetchone()
    c.execute(
        f"SELECT id,title,t.views,t.reactions,t.comments,t.shares,published_at,vk_post_id "
        f"FROM {TOTALS} WHERE {where} AND status='published' "
        "ORDER BY (t.views+t.reactions*3+t.comments*2+t.shares*4) DESC LIMIT 5",
        params,
    )
    top = c.fetchall()
    pl_stats = _platform_stats(c, f"{where} AND status='published'", params)

    return {
        "total_posts": total, "published": published, "scheduled": scheduled,
        "drafts": drafts, "on_review": on_review,
        "total_views": s["v"] or 0, "total_reactions": s["r"] or 0,
        "total_comments": s["cm"] or 0, "total_shares": s["sh"] or 0,
        "avg_views": _avg_views(_collected_views(s), published),
        "engagement_rate": _engagement(_collected_views(s), s["r"], s["cm"]),
        "top_posts": [dict(r) for r in top],
        "platform_stats": pl_stats,
    }


def _timeline(c, where: str, params: list, days: list) -> list[dict]:
    """
    Динамика по дням — одним запросом на весь период.

    Раньше здесь был цикл с отдельным SELECT на каждый день: неделя — 7
    обращений к базе, месяц — 30, квартал — 90, а экспорт за произвольный
    период мог попросить и больше. Запросы отличались только датой, то есть
    девяносто раз перечитывали одну и ту же выборку ради одной строки итога.

    Диапазон берётся сравнением с published_at напрямую (а не
    `published_at::date = ...`), чтобы условие ложилось на индекс
    `ix_posts_published`. Пустые дни SQL не вернёт — их дорисовываем здесь:
    графику нужен сплошной ряд, а «в этот день ничего не выходило» и «этого
    дня нет в ответе» для него разные вещи.
    """
    if not days:
        return []
    first = min(days).strftime("%Y-%m-%d")
    last = max(days).strftime("%Y-%m-%d")
    c.execute(
        f"SELECT published_at::date AS d, SUM(t.views) v, SUM(t.reactions) r, COUNT(*) p "
        f"FROM {TOTALS} WHERE {where} "
        "AND published_at >= %s::date AND published_at < %s::date + INTERVAL '1 day' "
        "GROUP BY 1",
        params + [first, last],
    )
    by_day = {row["d"].strftime("%Y-%m-%d"): row for row in c.fetchall()}

    result = []
    for day in days:
        ds = day.strftime("%Y-%m-%d")
        row = by_day.get(ds)
        result.append({
            "date": ds, "label": day.strftime("%d.%m"),
            "views": (row["v"] or 0) if row else 0,
            "reactions": (row["r"] or 0) if row else 0,
            "posts": (row["p"] or 0) if row else 0,
        })
    return result


def _period_days(period: str) -> list:
    days = {"week": 7, "month": 30, "quarter": 90}.get(period, 30)
    now = app_now()
    return [now - timedelta(days=i) for i in range(days - 1, -1, -1)]


def _export(c, where: str, params: list, start_date: str, end_date: str):
    try:
        dt_start = datetime.strptime(start_date, "%Y-%m-%d")
        dt_end = datetime.strptime(end_date, "%Y-%m-%d")
    except ValueError:
        raise HTTPException(status_code=400, detail="Формат дат: YYYY-MM-DD")

    end_next = (dt_end + timedelta(days=1)).strftime("%Y-%m-%d")
    start_str = dt_start.strftime("%Y-%m-%d")
    period = f"{where} AND published_at >= %s AND published_at < %s AND status='published'"
    period_params = params + [start_str, end_next]

    def count(sql: str, sql_params: list) -> int:
        c.execute(f"SELECT COUNT(*) FROM posts WHERE {sql}", sql_params)
        return c.fetchone()["count"]

    published = count(period, period_params)
    scheduled = count(f"{where} AND status='scheduled'", params)
    drafts = count(f"{where} AND status='draft'", params)
    total = count(f"{where} AND published_at >= %s AND published_at < %s",
                  params + [start_str, end_next])

    c.execute(
        f"SELECT SUM(t.views) v,SUM(t.reactions) r,SUM(t.comments) cm,SUM(t.shares) sh,"
        f"SUM(t.views_samples) vs "
        f"FROM {TOTALS} WHERE {period}",
        period_params,
    )
    s = c.fetchone()

    span = [dt_start + timedelta(days=i) for i in range((dt_end - dt_start).days + 1)]
    timeline = _timeline(c, where, params, span)
    pl_stats = _platform_stats(c, period, period_params, label_upper=True)

    c.execute(
        f"SELECT title,t.views,t.reactions,t.comments,t.shares,published_at "
        f"FROM {TOTALS} WHERE {period} "
        "ORDER BY (t.views+t.reactions*3+t.comments*2+t.shares*4) DESC LIMIT 10",
        period_params,
    )
    top_posts = c.fetchall()

    summary = {
        "total": total, "published": published, "scheduled": scheduled, "drafts": drafts,
        "total_views": s["v"] or 0, "total_reactions": s["r"] or 0,
        "total_comments": s["cm"] or 0, "total_shares": s["sh"] or 0,
        "avg_views": _avg_views(_collected_views(s), published),
        "eng": _engagement(_collected_views(s), s["r"], s["cm"]),
    }
    wb = _build_workbook(dt_start, dt_end, summary, timeline, pl_stats, top_posts)
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)

    month_ru = ["январь", "февраль", "март", "апрель", "май", "июнь",
                "июль", "август", "сентябрь", "октябрь", "ноябрь", "декабрь"]
    fname = f"аналитика_{month_ru[dt_start.month - 1]}_{dt_start.year}.xlsx"
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename*=UTF-8\'\'{quote(fname, safe='')}"},
    )


# ── Ручки: только выбор области, вся работа в ядре ──────────────────────────

@router.get("/api/analytics/summary")
def analytics_summary(user_id: int = Depends(get_current_user_id)):
    conn = get_db()
    try:
        scope, sp = posts_scope(user_id, conn, alias="")
        return _summary(conn.cursor(), scope, list(sp))
    finally:
        conn.close()


@router.get("/api/analytics/timeline")
def analytics_timeline(period: str = "month", user_id: int = Depends(get_current_user_id)):
    conn = get_db()
    try:
        scope, sp = posts_scope(user_id, conn, alias="")
        return _timeline(conn.cursor(), scope, list(sp), _period_days(period))
    finally:
        conn.close()


@router.get("/api/analytics/export")
def analytics_export(
    start_date: str = Query(...),
    end_date: str = Query(...),
    user_id: int = Depends(get_current_user_id),
):
    conn = get_db()
    try:
        scope, sp = posts_scope(user_id, conn, alias="")
        return _export(conn.cursor(), scope, list(sp), start_date, end_date)
    finally:
        conn.close()


@router.get("/api/analytics/report")
def analytics_report(
    start_date: str = Query(...),
    end_date: str = Query(...),
    org: str = Query("Молодёжный центр"),
    user_id: int = Depends(get_current_user_id),
):
    """Отчёт для учредителя по всем доступным пользователю постам."""
    conn = get_db()
    try:
        scope, sp = posts_scope(user_id, conn, alias="")
        return _report(conn.cursor(), scope, list(sp), org, start_date, end_date, None)
    finally:
        conn.close()


# ── Group-scoped Analytics ──────────────────────────────────────────────────

@router.get("/api/groups/{gid}/analytics/summary")
def group_analytics_summary(gid: int, user_id: int = Depends(get_current_user_id)):
    conn = get_db()
    try:
        require_group_member(gid, user_id, conn)
        return _summary(conn.cursor(), "group_id=%s", [gid])
    finally:
        conn.close()


@router.get("/api/groups/{gid}/analytics/timeline")
def group_analytics_timeline(gid: int, period: str = "month", user_id: int = Depends(get_current_user_id)):
    conn = get_db()
    try:
        require_group_member(gid, user_id, conn)
        return _timeline(conn.cursor(), "group_id=%s", [gid], _period_days(period))
    finally:
        conn.close()


@router.get("/api/groups/{gid}/analytics/report")
def group_analytics_report(
    gid: int,
    start_date: str = Query(...),
    end_date: str = Query(...),
    user_id: int = Depends(get_current_user_id),
):
    """
    Отчёт для учредителя. Название организации берём из названия группы —
    просить человека вводить его в поле было бы лишним: оно уже есть.
    """
    conn = get_db()
    try:
        require_group_member(gid, user_id, conn)
        c = conn.cursor()
        c.execute("SELECT name FROM groups WHERE id=%s", (gid,))
        row = c.fetchone()
        if not row:
            raise HTTPException(404, "Группа не найдена")
        return _report(c, "group_id=%s", [gid], row["name"], start_date, end_date, gid)
    finally:
        conn.close()


@router.get("/api/groups/{gid}/analytics/export")
def group_analytics_export(
    gid: int,
    start_date: str = Query(...),
    end_date: str = Query(...),
    user_id: int = Depends(get_current_user_id),
):
    conn = get_db()
    try:
        require_group_member(gid, user_id, conn)
        return _export(conn.cursor(), "group_id=%s", [gid], start_date, end_date)
    finally:
        conn.close()
